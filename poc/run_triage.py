"""執行分流：讀工單 → 規則＋本地模型 → 輸出 triage_result.csv 與 summary.md。

用法：python run_triage.py [工單檔 .xlsx 或 .csv] [--out 輸出資料夾]
triage_result.csv 是寫回工單系統的內部分流結果，結構化欄位用英文代碼；送到外部監控的只能是不含內容的 metadata。
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import llm
from rules import Ticket
from summary import write_summary
from triage import Triage, triage

HERE = Path(__file__).parent
FIELDS = ["ticket_id", "created_at", "channel", "customer_name", "subject", "body"]
OUTPUT_COLUMNS = ["ticket_id", "category", "route", "reason", "priority", "risk_flags",
                  "allowed_tools", "ai_draft", "next_step"]


def _to_datetime(value) -> datetime:
    return value if isinstance(value, datetime) else datetime.fromisoformat(str(value).strip())


def load_tickets(path: Path) -> list[Ticket]:
    """只讀題目定義的 6 個欄位。"""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            header = reader.fieldnames or []
            rows = list(reader)
    else:
        wb = load_workbook(path, read_only=True)
        values = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
        header = [str(h).strip() if h is not None else "" for h in values[0]]
        rows = [dict(zip(header, r)) for r in values[1:]]

    if missing := [f for f in FIELDS if f not in header]:
        raise ValueError(f"缺少欄位：{missing}")

    return [
        Ticket(
            ticket_id=str(row["ticket_id"]).strip(),
            created_at=_to_datetime(row["created_at"]),
            channel=str(row["channel"] or ""),
            customer_name=str(row["customer_name"] or ""),
            subject=str(row["subject"] or ""),
            body=str(row["body"] or ""),
        )
        for row in rows
        if row.get("ticket_id")
    ]


def write_csv(results: list[Triage], path: Path) -> None:
    # utf-8-sig：Excel 直接開啟 reason 裡的中文不亂碼；多值欄位以「;」分隔，沒有值就留空
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(OUTPUT_COLUMNS)
        for r in results:
            w.writerow([r.ticket_id, r.category, r.route, r.reason, r.priority, ";".join(r.risk_flags),
                        ";".join(r.allowed_tools), str(r.ai_draft).lower(), r.next_step])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", nargs="?", type=Path, default=HERE / "data" / "sample_tickets.xlsx")
    parser.add_argument("--out", type=Path, default=HERE / "output")
    args = parser.parse_args()

    try:
        llm.check_ready()
    except llm.ModelUnavailable as e:
        sys.exit(f"錯誤：{e}")

    tickets = load_tickets(args.input)
    results = []
    for i, t in enumerate(tickets, 1):
        try:
            results.append(triage(t))
        except llm.ModelUnavailable as e:
            sys.exit(f"錯誤：{e}")
        print(f"[{i}/{len(tickets)}] {t.ticket_id} → {results[-1].category}", flush=True)

    args.out.mkdir(parents=True, exist_ok=True)
    write_csv(results, args.out / "triage_result.csv")
    write_summary(results, tickets, args.input, args.out / "summary.md")

    routes = Counter(r.route for r in results)
    print(f"{len(results)} 筆 → {args.out / 'triage_result.csv'}（auto {routes['auto']}／human {routes['human']}）")
    print(f"摘要 → {args.out / 'summary.md'}")


if __name__ == "__main__":
    main()
